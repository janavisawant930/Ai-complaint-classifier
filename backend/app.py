import io
import os
import re
from datetime import datetime
from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from werkzeug.security import generate_password_hash, check_password_hash
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import Config
from database import init_db, query_db, execute_db, get_db_connection
from classifier import classify_complaint
from rag_service import classify_with_rag, build_or_get_vectorstore

app = Flask(__name__)
app.config.from_object(Config)

# Enable CORS for frontend integration
if Config.CORS_ORIGINS == "*":
    CORS(app, resources={r"/api/*": {"origins": "*"}}, supports_credentials=True)
else:
    origins = [o.strip() for o in Config.CORS_ORIGINS.split(",") if o.strip()]
    CORS(app, resources={r"/api/*": {"origins": origins}}, supports_credentials=True)

# Helper function to format complaint output for API
def format_complaint(c):
    if not c:
        return None
    # Keep complaint and description synced for compatibility
    desc = c.get("description") or c.get("complaint") or ""
    return {
        "id": c.get("id"),
        "user": c.get("user"),
        "user_id": c.get("user_id"),
        "title": c.get("title") or (desc[:40] + "..." if len(desc) > 40 else desc),
        "description": desc,
        "complaint": desc,  # frontend compatibility alias
        "category": c.get("category"),
        "priority": c.get("priority"),
        "confidence": int(c.get("confidence") or 85),
        "status": c.get("status") or "Pending",
        "attachment": c.get("attachment"),
        "date": c.get("date") or datetime.now().strftime("%d %b %Y"),
        "created_at": str(c.get("created_at")) if c.get("created_at") else None
    }

# ----------------------------------------------------
# Health and Info Endpoints
# ----------------------------------------------------
@app.route("/", methods=["GET"])
def index():
    return jsonify({
        "name": "ComplainAI Backend API",
        "status": "online",
        "version": "1.0.0",
        "endpoints": {
            "auth_signup": "POST /api/auth/signup",
            "auth_signin": "POST /api/auth/signin",
            "get_complaints": "GET /api/complaints",
            "export_complaints": "GET /api/complaints/export",
            "submit_complaint": "POST /api/complaints",
            "get_complaint": "GET /api/complaints/<id>",
            "update_status": "PATCH /api/complaints/<id>/status",
            "ai_classification": "POST /api/classify",
            "analytics": "GET /api/analytics"
        }
    })

@app.route("/api/health", methods=["GET"])
def health_check():
    db_status = "connected"
    try:
        query_db("SELECT 1", one=True)
    except Exception as e:
        db_status = f"disconnected: {str(e)}"

    rag_status = "ready" if Config.GEMINI_API_KEY else "unconfigured (missing GEMINI_API_KEY)"

    return jsonify({
        "status": "healthy" if db_status == "connected" else "degraded",
        "database": db_status,
        "rag_pipeline": rag_status,
        "gemini_model": Config.GEMINI_MODEL,
        "timestamp": datetime.now().isoformat()
    })


# ----------------------------------------------------
# 1. User Sign Up
# ----------------------------------------------------
@app.route("/api/auth/signup", methods=["POST"])
@app.route("/api/signup", methods=["POST"])
def signup():
    data = request.get_json() or {}
    
    name = (data.get("fullName") or data.get("name") or "").strip()
    email = (data.get("email") or "").strip().lower()
    username = (data.get("username") or "").strip().lower()
    password = data.get("password") or ""
    confirm_password = data.get("confirmPassword") or password

    # Validations
    if not name or len(name) < 2:
        return jsonify({"success": False, "error": "Full Name must be at least 2 characters."}), 400

    if not email or "@" not in email:
        return jsonify({"success": False, "error": "A valid email address is required."}), 400

    if not username or len(username) < 3:
        return jsonify({"success": False, "error": "Username must be at least 3 characters."}), 400

    if not password or len(password) < 4:
        return jsonify({"success": False, "error": "Password must be at least 4 characters."}), 400

    if password != confirm_password:
        return jsonify({"success": False, "error": "Passwords do not match."}), 400

    try:
        # Check if email exists
        existing_email = query_db("SELECT id FROM users WHERE LOWER(email) = %s", (email,), one=True)
        if existing_email:
            return jsonify({"success": False, "error": "An account with this email address already exists."}), 409

        # Check if username exists
        existing_user = query_db("SELECT id FROM users WHERE LOWER(username) = %s", (username,), one=True)
        if existing_user:
            return jsonify({"success": False, "error": "This username is already taken. Please choose another."}), 409

        # Generate avatar initials
        name_parts = [p for p in name.split() if p]
        initials = (name_parts[0][0] + name_parts[1][0]).upper() if len(name_parts) >= 2 else name[:2].upper()

        password_hash = generate_password_hash(password)
        role = data.get("role") or "Client / User"

        user_id = execute_db("""
            INSERT INTO users (username, email, password, name, role, initials)
            VALUES (%s, %s, %s, %s, %s, %s)
        """, (username, email, password_hash, name, role, initials))

        return jsonify({
            "success": True,
            "message": "Account created successfully! You can now sign in.",
            "user": {
                "id": user_id,
                "username": username,
                "email": email,
                "name": name,
                "role": role,
                "initials": initials
            }
        }), 201

    except Exception as e:
        return jsonify({"success": False, "error": f"Database error: {str(e)}"}), 500

# ----------------------------------------------------
# 2. User Sign In
# ----------------------------------------------------
@app.route("/api/auth/signin", methods=["POST"])
@app.route("/api/auth/login", methods=["POST"])
@app.route("/api/signin", methods=["POST"])
@app.route("/api/login", methods=["POST"])
def signin():
    data = request.get_json() or {}
    identifier = (data.get("identifier") or data.get("email") or data.get("username") or "").strip()
    password = data.get("password") or ""

    if not identifier or not password:
        return jsonify({"success": False, "error": "Please provide both identifier (email/username) and password."}), 400

    try:
        user = query_db("""
            SELECT id, username, email, password, name, role, initials
            FROM users
            WHERE LOWER(email) = %s OR LOWER(username) = %s
        """, (identifier.lower(), identifier.lower()), one=True)

        if not user:
            return jsonify({"success": False, "error": "Invalid credentials. User not found."}), 401

        # Check password with hash support or fallback to plain text for legacy/demo seeds
        stored_pw = user["password"]
        password_valid = False
        if stored_pw.startswith("pbkdf2:") or stored_pw.startswith("scrypt:"):
            password_valid = check_password_hash(stored_pw, password)
        else:
            password_valid = (stored_pw == password)

        if not password_valid:
            return jsonify({"success": False, "error": "Invalid credentials. Password is incorrect."}), 401

        return jsonify({
            "success": True,
            "message": f"Welcome back, {user['name']}!",
            "user": {
                "id": user["id"],
                "username": user["username"],
                "email": user["email"],
                "name": user["name"],
                "role": user["role"] or "Client / User",
                "initials": user["initials"] or user["name"][:2].upper()
            }
        }), 200

    except Exception as e:
        return jsonify({"success": False, "error": f"Database error: {str(e)}"}), 500

# ----------------------------------------------------
# 3. AI Classification Endpoint (LangChain + RAG)
# ----------------------------------------------------
@app.route("/api/classify", methods=["POST"])
@app.route("/api/ai/classify", methods=["POST"])
def classify():
    data = request.get_json() or {}
    text = data.get("text") or data.get("description") or data.get("complaint") or ""
    title = data.get("title") or ""
    combined_text = f"{title} {text}".strip()

    if not combined_text:
        return jsonify({"success": False, "error": "No complaint text provided for classification."}), 400

    ai_result = classify_with_rag(text, title=title)
    return jsonify({
        "success": True,
        "category": ai_result["category"],
        "priority": ai_result["priority"],
        "confidence": ai_result["confidence"],
        "reason": ai_result.get("reason", ""),
        "suggested_resolution": ai_result.get("suggested_resolution", ""),
        "matched_keywords": ai_result.get("matched_keywords", []),
        "retrieved_context": ai_result.get("retrieved_context", [])
    }), 200

# ----------------------------------------------------
# 4. Submit Complaint (with LangChain + RAG Classification)
# ----------------------------------------------------
@app.route("/api/complaints", methods=["POST"])
@app.route("/api/complaint", methods=["POST"])
def submit_complaint():
    data = request.get_json() or {}

    title = (data.get("title") or "").strip()
    description = (data.get("description") or data.get("complaint") or "").strip()
    user_name = (data.get("user") or data.get("userName") or "Anonymous").strip()
    user_id = data.get("user_id")
    attachment = data.get("attachment")

    if not title and not description:
        return jsonify({"success": False, "error": "Complaint title and description are required."}), 400

    if not title:
        title = description[:50] + "..." if len(description) > 50 else description
    if not description:
        description = title

    # Run LangChain + RAG AI Classification
    category = data.get("category")
    priority = data.get("priority")
    confidence = data.get("confidence")
    ai_reason = None
    ai_resolution = None

    if not category or not priority or confidence is None:
        ai_result = classify_with_rag(description, title=title)
        category = ai_result["category"]
        priority = ai_result["priority"]
        confidence = ai_result["confidence"]
        ai_reason = ai_result.get("reason")
        ai_resolution = ai_result.get("suggested_resolution")
    else:
        try:
            confidence = int(confidence)
        except (ValueError, TypeError):
            confidence = 90

    status = data.get("status") or "Pending"
    date_str = data.get("date") or datetime.now().strftime("%d %b %Y")

    try:
        # Generate Complaint ID (e.g. CMP-1007)
        last_complaint = query_db("""
            SELECT id FROM complaints
            WHERE id LIKE %s
            ORDER BY created_at DESC, id DESC
            LIMIT 1
        """, ('CMP-%',), one=True)

        new_id = None
        if last_complaint and last_complaint["id"]:
            match = re.search(r'CMP-(\d+)', last_complaint["id"])
            if match:
                next_num = int(match.group(1)) + 1
                new_id = f"CMP-{next_num}"

        if not new_id:
            count_res = query_db("SELECT COUNT(*) AS total FROM complaints", one=True)
            total_count = (count_res["total"] if count_res else 0) + 1001
            new_id = f"CMP-{total_count}"

        execute_db("""
            INSERT INTO complaints (id, user, user_id, title, description, category, priority, confidence, status, attachment, date)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        """, (new_id, user_name, user_id, title, description, category, priority, confidence, status, attachment, date_str))

        created_complaint = query_db("SELECT * FROM complaints WHERE id = %s", (new_id,), one=True)
        formatted = format_complaint(created_complaint)

        return jsonify({
            "success": True,
            "message": f"Complaint {new_id} submitted and classified as {category} ({priority} Priority)!",
            "complaint": formatted,
            "ai_result": {
                "category": category,
                "priority": priority,
                "confidence": confidence,
                "reason": ai_reason,
                "suggested_resolution": ai_resolution
            }
        }), 201

    except Exception as e:
        return jsonify({"success": False, "error": f"Failed to submit complaint: {str(e)}"}), 500


# ----------------------------------------------------
# 5. Get All Complaints (with search, filters & sort)
# ----------------------------------------------------
@app.route("/api/complaints", methods=["GET"])
def get_complaints():
    category = request.args.get("category")
    status = request.args.get("status")
    priority = request.args.get("priority")
    search = request.args.get("search", "").strip()
    sort_by = request.args.get("sort", "newest")

    try:
        sql = "SELECT * FROM complaints WHERE 1=1"
        params = []

        if category and category != "All":
            sql += " AND category = %s"
            params.append(category)

        if status and status != "All":
            sql += " AND status = %s"
            params.append(status)

        if priority and priority != "All":
            sql += " AND priority = %s"
            params.append(priority)

        if search:
            search_pattern = f"%{search}%"
            sql += " AND (id LIKE %s OR user LIKE %s OR title LIKE %s OR description LIKE %s OR category LIKE %s OR priority LIKE %s)"
            params.extend([search_pattern] * 6)

        if sort_by == "confidence-high":
            sql += " ORDER BY confidence DESC"
        elif sort_by == "confidence-low":
            sql += " ORDER BY confidence ASC"
        elif sort_by == "oldest":
            sql += " ORDER BY created_at ASC, id ASC"
        else:
            # Default newest
            sql += " ORDER BY created_at DESC, id DESC"

        complaints_raw = query_db(sql, params)
        formatted = [format_complaint(c) for c in complaints_raw]

        return jsonify({
            "success": True,
            "count": len(formatted),
            "complaints": formatted
        }), 200

    except Exception as e:
        return jsonify({"success": False, "error": f"Failed to fetch complaints: {str(e)}"}), 500

# ----------------------------------------------------
# 5b. Export Complaints to Excel (.xlsx) using OpenPyXL
# ----------------------------------------------------
@app.route("/api/complaints/export", methods=["GET"])
@app.route("/api/export", methods=["GET"])
def export_complaints():
    category = request.args.get("category")
    status = request.args.get("status")
    priority = request.args.get("priority")
    search = request.args.get("search", "").strip()
    sort_by = request.args.get("sort", "newest")

    try:
        sql = "SELECT * FROM complaints WHERE 1=1"
        params = []

        if category and category != "All":
            sql += " AND category = %s"
            params.append(category)

        if status and status != "All":
            sql += " AND status = %s"
            params.append(status)

        if priority and priority != "All":
            sql += " AND priority = %s"
            params.append(priority)

        if search:
            search_pattern = f"%{search}%"
            sql += " AND (id LIKE %s OR user LIKE %s OR title LIKE %s OR description LIKE %s OR category LIKE %s OR priority LIKE %s)"
            params.extend([search_pattern] * 6)

        if sort_by == "confidence-high":
            sql += " ORDER BY confidence DESC"
        elif sort_by == "confidence-low":
            sql += " ORDER BY confidence ASC"
        elif sort_by == "oldest":
            sql += " ORDER BY created_at ASC, id ASC"
        else:
            # Default newest
            sql += " ORDER BY created_at DESC, id DESC"

        complaints_raw = query_db(sql, params) or []

        # Initialize OpenPyXL workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Complaints Report"

        # Ensure Excel gridlines are shown
        ws.views.sheetView[0].showGridLines = True

        # Header list
        headers = [
            "Complaint ID",
            "User",
            "Complaint Title",
            "Complaint Description",
            "AI Category",
            "Priority",
            "Confidence (%)",
            "Status",
            "Date",
            "Created At"
        ]

        # Professional styling definitions
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        thin_border_side = Side(border_style="thin", color="D1D5DB")
        border = Border(
            left=thin_border_side,
            right=thin_border_side,
            top=thin_border_side,
            bottom=thin_border_side
        )

        # Write header row
        ws.append(headers)
        ws.row_dimensions[1].height = 28

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border

        # Fonts & alignments for data rows
        row_font = Font(name="Calibri", size=10)
        align_left = Alignment(horizontal="left", vertical="center")
        align_center = Alignment(horizontal="center", vertical="center")
        align_desc = Alignment(horizontal="left", vertical="top", wrap_text=True)

        # Write complaint rows
        for row_idx, c in enumerate(complaints_raw, start=2):
            desc = c.get("description") or c.get("complaint") or ""
            title = c.get("title") or (desc[:40] + "..." if len(desc) > 40 else desc)
            conf = int(c.get("confidence") or 0)
            created_at_val = str(c.get("created_at")) if c.get("created_at") else ""

            row_data = [
                c.get("id") or "",
                c.get("user") or "",
                title,
                desc,
                c.get("category") or "",
                c.get("priority") or "",
                conf,
                c.get("status") or "Pending",
                c.get("date") or "",
                created_at_val
            ]
            ws.append(row_data)
            ws.row_dimensions[row_idx].height = 22

            # Apply cell styles
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = row_font
                cell.border = border

                if col_idx in (1, 5, 6, 7, 8, 9, 10):
                    cell.alignment = align_center
                elif col_idx == 4:
                    cell.alignment = align_desc
                else:
                    cell.alignment = align_left

        # Adjust column widths based on contents with sensible limits
        col_width_constraints = {
            1: (15, 20),   # ID
            2: (16, 26),   # User
            3: (22, 36),   # Title
            4: (30, 55),   # Description
            5: (18, 26),   # Category
            6: (12, 18),   # Priority
            7: (15, 18),   # Confidence
            8: (14, 18),   # Status
            9: (14, 20),   # Date
            10: (20, 24),  # Created At
        }

        for col_idx, col in enumerate(ws.columns, start=1):
            min_w, max_w = col_width_constraints.get(col_idx, (12, 30))
            max_len = 0
            for cell in col:
                val_str = str(cell.value or "")
                if "\n" in val_str:
                    lines = val_str.split("\n")
                    max_len = max(max_len, max(len(l) for l in lines))
                else:
                    max_len = max(max_len, len(val_str))
            calculated_width = max(min_w, min(max_len + 3, max_w))
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = calculated_width

        # Save to in-memory bytes stream
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = "complaints_report.xlsx"
        return send_file(
            output,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=filename
        )

    except Exception as e:
        return jsonify({"success": False, "error": f"Failed to export complaints to Excel: {str(e)}"}), 500

# ----------------------------------------------------
# 6. Get Single Complaint
# ----------------------------------------------------
@app.route("/api/complaints/<complaint_id>", methods=["GET"])
def get_single_complaint(complaint_id):
    try:
        complaint = query_db("SELECT * FROM complaints WHERE id = %s", (complaint_id,), one=True)
        if not complaint:
            return jsonify({"success": False, "error": f"Complaint with ID '{complaint_id}' not found."}), 404

        return jsonify({
            "success": True,
            "complaint": format_complaint(complaint)
        }), 200

    except Exception as e:
        return jsonify({"success": False, "error": f"Failed to retrieve complaint: {str(e)}"}), 500

# ----------------------------------------------------
# 7. Update Complaint Status
# ----------------------------------------------------
@app.route("/api/complaints/<complaint_id>/status", methods=["PATCH", "PUT"])
@app.route("/api/complaints/<complaint_id>", methods=["PATCH", "PUT"])
def update_complaint_status(complaint_id):
    data = request.get_json() or {}
    new_status = data.get("status")

    valid_statuses = ["Pending", "In Progress", "Resolved"]
    if not new_status or new_status not in valid_statuses:
        return jsonify({
            "success": False,
            "error": f"Invalid status. Must be one of: {', '.join(valid_statuses)}"
        }), 400

    try:
        complaint = query_db("SELECT id FROM complaints WHERE id = %s", (complaint_id,), one=True)
        if not complaint:
            return jsonify({"success": False, "error": f"Complaint with ID '{complaint_id}' not found."}), 404

        execute_db("UPDATE complaints SET status = %s WHERE id = %s", (new_status, complaint_id))
        updated = query_db("SELECT * FROM complaints WHERE id = %s", (complaint_id,), one=True)

        return jsonify({
            "success": True,
            "message": f"Complaint {complaint_id} status updated to {new_status}.",
            "complaint": format_complaint(updated)
        }), 200

    except Exception as e:
        return jsonify({"success": False, "error": f"Failed to update status: {str(e)}"}), 500

# ----------------------------------------------------
# 8. Analytics Endpoint
# ----------------------------------------------------
@app.route("/api/analytics", methods=["GET"])
def get_analytics():
    try:
        complaints_raw = query_db("SELECT * FROM complaints ORDER BY created_at DESC")
        total = len(complaints_raw)

        categories_list = [
            "Billing",
            "Technical Issues",
            "Product/Service",
            "Account",
            "Delivery",
            "Fraud/Security",
            "Other"
        ]

        # Category distribution
        by_category = {}
        for cat in categories_list:
            count = sum(1 for c in complaints_raw if c.get("category") == cat)
            by_category[cat] = {
                "count": count,
                "percentage": round((count / total) * 100, 1) if total > 0 else 0
            }

        # Status distribution
        pending = sum(1 for c in complaints_raw if c.get("status") == "Pending")
        in_progress = sum(1 for c in complaints_raw if c.get("status") == "In Progress")
        resolved = sum(1 for c in complaints_raw if c.get("status") == "Resolved")

        # Priority distribution
        high_priority = sum(1 for c in complaints_raw if c.get("priority") == "High")
        medium_priority = sum(1 for c in complaints_raw if c.get("priority") == "Medium")
        low_priority = sum(1 for c in complaints_raw if c.get("priority") == "Low")

        # Confidence calculation
        avg_confidence = (
            round(sum(c.get("confidence", 0) for c in complaints_raw) / total)
            if total > 0 else 0
        )

        resolution_rate = round((resolved / total) * 100, 1) if total > 0 else 0

        recent_complaints = [format_complaint(c) for c in complaints_raw[:5]]

        return jsonify({
            "success": True,
            "data": {
                "total": total,
                "pending": pending,
                "resolved": resolved,
                "in_progress": in_progress,
                "high_priority": high_priority,
                "medium_priority": medium_priority,
                "low_priority": low_priority,
                "avg_confidence": avg_confidence,
                "resolution_rate": f"{resolution_rate}%",
                "by_category": by_category,
                "by_priority": {
                    "High": high_priority,
                    "Medium": medium_priority,
                    "Low": low_priority
                },
                "by_status": {
                    "Pending": pending,
                    "In Progress": in_progress,
                    "Resolved": resolved
                },
                "recent_complaints": recent_complaints
            }
        }), 200

    except Exception as e:
        return jsonify({"success": False, "error": f"Failed to compute analytics: {str(e)}"}), 500

# Error Handlers
@app.errorhandler(404)
def not_found_error(e):
    return jsonify({"success": False, "error": "API route not found."}), 404

@app.errorhandler(405)
def method_not_allowed(e):
    return jsonify({"success": False, "error": "HTTP method not allowed for this route."}), 405

@app.errorhandler(500)
def internal_error(e):
    return jsonify({"success": False, "error": "Internal server error occurred."}), 500

if __name__ == "__main__":
    # Initialize MySQL tables on startup
    try:
        init_db()
    except Exception as e:
        print(f"[Warning] Database initialization notice: {e}")

    # Warm-up / verify RAG Chroma Vector Database on startup
    try:
        if Config.GEMINI_API_KEY:
            build_or_get_vectorstore()
            print("[RAG] Chroma Vector Database loaded and ready.")
        else:
            print("[RAG Notice] GEMINI_API_KEY not found in environment; running in fallback mode.")
    except Exception as e:
        print(f"[Warning] RAG vectorstore warmup notice: {e}")

    print(f" * Starting ComplainAI Flask Backend on http://127.0.0.1:{Config.PORT}")
    app.run(host="0.0.0.0", port=Config.PORT, debug=Config.DEBUG)

